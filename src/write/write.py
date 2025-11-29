import argparse
import copy
import csv
import json
import os
from lxml import etree
import conllu
from openpyxl import load_workbook

from src.create_tei import construct_sentence_from_list, \
    construct_paragraph_from_list, TeiDocument, build_tei_etrees, build_links, build_complete_tei, convert_bibl
from src.validate.validate_xml import validate_xml


def form_paragraphs(annotated_source_divs, metadata):
    etree_source_divs = []
    div_name = "#DIV_NAME#"
    for div_i, div_tuple in enumerate(annotated_source_divs):
        div_name, div = div_tuple
        if div_name[:-1] not in metadata:
            raise Exception("Not in metadata:", div_name[:-1])
        div_metadata = metadata[div_name[:-1]]

        etree_source_paragraphs = []

        for par_i, paragraph_tuple in enumerate(div):
            par_name, paragraph = paragraph_tuple
            etree_source_sentences = []

            for sentence_id, sentence in enumerate(paragraph):
                if len(sentence) > 0:
                    conllu_parsed = conllu.parse(sentence)[0]
                    etree_source_sentences.append(
                        construct_sentence_from_list(str(sentence_id + 1), conllu_parsed, True))

            etree_source_paragraphs.append(construct_paragraph_from_list(div_name, par_name, etree_source_sentences))

        etree_source_divs.append((etree_source_paragraphs, div_name, div_metadata))

    return etree_source_divs, div_name


def read_metadata(args, delimiter='|'):
    texts_metadata = []
    with open(args.texts_metadata, 'r', encoding='utf-8-sig') as file:
        csvreader = csv.reader(file, delimiter=delimiter, quotechar='"')
        column_names = []
        for i, row in enumerate(csvreader):
            if i == 0:
                column_names = row
                continue
            else:
                row_dict = {}
                for j, content in enumerate(row):
                    if j in range(len(column_names) - 1): # and content.strip():
                        row_dict[column_names[j]] = content.strip()
                    #else:
                    #    print('fail', row)
                    #    print(j, column_names)
                texts_metadata.append(row_dict)

    # handle teachers
    teachers_metadata = {}
    with open(args.teachers_metadata, 'r', encoding='utf-8-sig') as file:
        csvreader = csv.reader(file, delimiter='\t', quotechar='"')
        column_names = []
        for i, row in enumerate(csvreader):
            if i == 0:
                column_names = row
                continue
            else:
                row_dict = {}
                for j, content in enumerate(row):
                    row_dict[column_names[j]] = content
                row_dict['Ime in priimek'] = row_dict['Ime in priimek'].strip()
                teachers_metadata[row_dict['Ime in priimek']] = row_dict

    # handle authors
    authors_metadata = {}
    with open(args.authors_metadata, 'r', encoding='utf-8-sig') as file:
        csvreader = csv.reader(file, delimiter=delimiter, quotechar='"')
        #csvreader = csv.reader(file, delimiter=',', quotechar='"')
        column_names = []
        for i, row in enumerate(csvreader):
            if i == 0:
                column_names = row
                continue
            elif i == 1:
                active_column_name = ''
                for j, sub_name in enumerate(row):
                    if column_names[j]:
                        active_column_name = column_names[j]
                    if sub_name:
                        column_names[j] = f'{active_column_name} - {sub_name}'
                continue
            elif i == 2:
                continue
            else:
                row_dict = {}
                for j, content in enumerate(row):
                    row_dict[column_names[j]] = content.strip()
                row_dict['Ime in priimek'] = row_dict['Ime in priimek'].strip()
                authors_metadata[row_dict['Ime in priimek']] = row_dict

    translations = {}
    with open(args.translations, 'r', encoding='utf-8-sig') as file:
        csvreader = csv.reader(file, delimiter='\t', quotechar='"')
        for row in csvreader:
            translations[row[0]] = row[1]

    return texts_metadata, authors_metadata, teachers_metadata, translations


def cell_to_str(cell):
    if cell is None or cell.value is None:
        return ''
    elif cell.data_type == 's':
        return cell.value
    if cell.data_type == 'n':
        return str(cell.value)
    elif cell.data_type == 'd':
        return cell.value.strftime('%d/%m/%Y')
    else:
        raise Exception(f'Error converting cell value to string: {cell.value}')

def read_metadata_from_excel(args):
    excel = load_workbook(args.metadata_excel, data_only=True)

    # parse texts sheet
    texts_sheet = excel['L+ besedila']
    texts_metadata = []
    column_names = []

    for i, row in enumerate(texts_sheet.iter_rows(values_only=False)):
        if i == 0:
            column_names = [cell_to_str(cell) for cell in row]
            continue
        elif not cell_to_str(row[0]):
            continue
        else:
            row_dict = {}
            for j, content in enumerate(row):
                if j in range(len(column_names) - 1):
                    row_dict[column_names[j]] = cell_to_str(content)
            texts_metadata.append(row_dict)

    # parse teachers sheet
    teachers_sheet = excel['Kode učiteljev']
    teachers_metadata = {}
    column_names = []
    for i, row in enumerate(teachers_sheet.iter_rows(values_only=False)):
        if i == 0:
            column_names = [cell_to_str(cell) for cell in row]
            continue
        elif not cell_to_str(row[0]):
            continue
        else:
            row_dict = {}
            for j, content in enumerate(row):
                row_dict[column_names[j]] = cell_to_str(content)
            row_dict['Ime in priimek'] = row_dict['Ime in priimek'].strip()
            teachers_metadata[row_dict['Ime in priimek']] = row_dict

    # parse authors sheet
    authors_sheet = excel['L+ tvorci']
    authors_metadata = {}
    column_names = []
    for i, row in enumerate(authors_sheet.iter_rows(values_only=False)):
        if i == 0:
            column_names = [cell_to_str(cell) for cell in row]
            continue
        elif i == 1:
            active_column_name = ''
            for j, sub_name in enumerate(row):
                sub_name = cell_to_str(sub_name)
                if column_names[j]:
                    active_column_name = column_names[j]
                if sub_name:
                    column_names[j] = f'{active_column_name} - {sub_name}'
            continue
        elif i == 2:
            continue
        elif not cell_to_str(row[0]):
            continue
        else:
            row_dict = {}
            for j, content in enumerate(row):
                row_dict[column_names[j]] = cell_to_str(content)
            row_dict['Ime in priimek'] = row_dict['Ime in priimek'].strip()
            authors_metadata[row_dict['Ime in priimek']] = row_dict

    # parse translations sheet
    translations_sheet = excel['Prevod kategorij v ang.']
    translations = {}
    for i, row in enumerate(translations_sheet.iter_rows(values_only=False)):
        if i < 3:
            continue
        if cell_to_str(row[0]):
            translations[cell_to_str(row[0])] = cell_to_str(row[1])

    return texts_metadata, authors_metadata, teachers_metadata, translations


def process_metadata(args):
    texts_metadata, authors_metadata, teachers_metadata, translations = read_metadata_from_excel(args)

    metadata = {}
    for document_metadata in texts_metadata:
        if not document_metadata:
            print(document_metadata)
            continue
        document_metadata['Tvorec'] = document_metadata['Tvorec'].strip()
        if document_metadata['Tvorec'] not in authors_metadata:
            if document_metadata['Tvorec']:
                print("Missing author data:", document_metadata['Tvorec'], document_metadata)
            continue
        author_metadata = authors_metadata[document_metadata['Tvorec']]
        metadata_el = {}
        for attribute_name_sl, attribute_name_en in translations.items():
            if attribute_name_sl in document_metadata:
                if attribute_name_sl == 'Ocena':
                    grade = f'{document_metadata[attribute_name_sl]} od {document_metadata["Najvišja možna ocena"]}' if document_metadata[attribute_name_sl] and document_metadata["Najvišja možna ocena"] else ''
                    metadata_el[attribute_name_en] = grade
                elif attribute_name_sl == 'Tvorec':
                    metadata_el[attribute_name_en] = author_metadata['Koda tvorca']
                elif attribute_name_sl == 'Učitelj':
                    metadata_el[attribute_name_en] = teachers_metadata[document_metadata['Učitelj']]['Koda'] if document_metadata['Učitelj'] in teachers_metadata else None
                else:
                    metadata_el[attribute_name_en] = document_metadata[attribute_name_sl]
            elif attribute_name_sl in author_metadata:
                metadata_el[attribute_name_en] = author_metadata[attribute_name_sl]
            elif attribute_name_sl == 'Ime šole, Fakulteta':
                curr_school = []
                if author_metadata["Trenutno šolanje - Ime šole"]:
                    curr_school.append(author_metadata["Trenutno šolanje - Ime šole"])
                if author_metadata["Trenutno šolanje - Fakulteta"]:
                    curr_school.append(author_metadata["Trenutno šolanje - Fakulteta"])
                metadata_el['    '] = ', '.join(curr_school)
            elif attribute_name_sl == 'Stopnja študija':
                metadata_el[attribute_name_en] = author_metadata['Trenutno šolanje - Stopnja študija']
            elif attribute_name_sl == 'Leto študija':
                metadata_el[attribute_name_en] = author_metadata['Trenutno šolanje - Leto študija']
            elif attribute_name_sl == 'Ostali jeziki':
                metadata_el[attribute_name_en] = ','.join([k[16:] for k, v in author_metadata.items() if k[:13] == 'Ostali jeziki' and v == 'ja'])
            elif attribute_name_sl == 'Kje učenje':
                metadata_el[attribute_name_en] = author_metadata['Življenje v Sloveniji pred tem programom - Kje?']
            elif attribute_name_sl == 'Koliko časa učenje?':
                metadata_el[attribute_name_en] = author_metadata['Življenje v Sloveniji pred tem programom - Koliko časa?']
            elif attribute_name_sl == 'Učbeniki':
                metadata_el[attribute_name_en] = author_metadata['Učenje slovenščine pred tem programom - Učbeniki']
            elif attribute_name_sl == 'Kje?':
                metadata_el[attribute_name_en] = author_metadata['Učenje slovenščine pred L+ - Kje?']
            elif attribute_name_sl == 'Koliko časa?':
                metadata_el[attribute_name_en] = author_metadata['Učenje slovenščine pred L+ - Koliko čas?']
            else:
                raise Exception(f'{attribute_name_sl} not found!')

        metadata[metadata_el['Text ID']] = metadata_el

    return metadata

def write_tei(annotated_source_divs, annotated_target_divs, document_edges, args):
    print('BUILDING LINKS...')
    etree_links = build_links(document_edges)

    with open(os.path.join(args.results_folder, f"links.xml"), 'w', encoding='utf-8') as tf:
        tf.write(etree.tostring(etree_links, pretty_print=True, encoding='utf-8').decode())

    with open(os.path.join(args.results_folder, f"links.json"), 'w', encoding='utf-8') as jf:
        json.dump(document_edges, jf, ensure_ascii=False, indent="  ")

    print('WRITTING TEI...')
    etree_source_documents = []
    etree_target_documents = []

    print('PREPARING METADATA FOR BIBL...')
    metadata = process_metadata(args)

    print('WRITING SOURCE FILES...')
    etree_source_divs, source_div_name = form_paragraphs(annotated_source_divs, metadata)

    print('WRITING TARGET FILES...')
    etree_target_divs, target_div_name = form_paragraphs(annotated_target_divs, metadata)

    print('APPENDING DOCUMENT...')
    etree_source_documents.append(
        TeiDocument(source_div_name,
                    etree_source_divs, etree_target_divs))
    etree_target_documents.append(
        TeiDocument(target_div_name,
                    etree_target_divs, etree_source_divs))

    print('BUILDING TEI DOCUMENTS...')
    etree_source = build_tei_etrees(etree_source_documents)
    etree_target = build_tei_etrees(etree_target_documents)

    # to reduce RAM usage you may process the following in two steps, firstly write all but complete (by commenting complete tree code), secondly write only complete (by commenting "Writting all but complete" section of code and "deepcopy" function)
    print('Writting all but complete')
    with open(os.path.join(args.results_folder, f"source.xml"), 'w', encoding='utf-8') as sf:
        sf.write(etree.tostring(etree_source[0], pretty_print=True, encoding='utf-8').decode())

    with open(os.path.join(args.results_folder, f"target.xml"), 'w', encoding='utf-8') as tf:
        tf.write(etree.tostring(etree_target[0], pretty_print=True, encoding='utf-8').decode())

    print('COMPLETE TREE CREATION...')
    complete_etree = build_complete_tei(copy.deepcopy(etree_source), copy.deepcopy(etree_target), etree_links)
    # complete_etree = build_complete_tei(etree_source, etree_target, etree_links)

    print('WRITING COMPLETE TREE')
    complete_filepath = os.path.join(args.results_folder, f"complete.xml")
    with open(complete_filepath, 'w', encoding='utf-8') as tf:
        tf.write(etree.tostring(complete_etree, pretty_print=True, encoding='utf-8').decode())

    print('VALIDATING TEI FILE')
    if validate_xml(complete_filepath):
        # check if all links have matching text
        print('SUCCESS:', complete_filepath)
    else:
        print('INVALID TEI FILE:', complete_filepath)


if __name__ == '__main__':
    parser = argparse.ArgumentParser(
        description='Merges svala data, raw data and metadata into TEI format (useful for corpora like KOST).')
    parser.add_argument('--metadata_excel', default='D:\projects\KOST\svala-scripts/data/KOST test/KOST 2.0, 25-08.xlsm',
                        help='KOST metadata location')
    args = parser.parse_args()

    read_metadata_from_excel(args)